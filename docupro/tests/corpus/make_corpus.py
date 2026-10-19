"""Generate the DOCUPRO test corpus deterministically.

Run: python tests/corpus/make_corpus.py <output_dir>
The corpus is generated rather than committed so the repository stays small
and every fixture is reproducible.
"""

from __future__ import annotations

import io
import os
import random
import sys
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw, ImageFont

random.seed(42)

FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
]


def font(size: int = 22):
    for path in FONT_CANDIDATES:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def chart_image(width: int = 640, height: int = 360) -> Image.Image:
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    draw.rectangle([40, 40, width - 40, height - 60], outline="#cbd5e1")
    bars = [("Jan", 120), ("Fév", 180), ("Mar", 90), ("Avr", 210)]
    for index, (label, value) in enumerate(bars):
        x = 90 + index * 120
        draw.rectangle([x, height - 90 - value, x + 60, height - 90], fill="#2563eb")
        draw.text((x + 8, height - 78), label, fill="#0f172a", font=font(16))
    draw.text((60, 55), "Ventes trimestrielles", fill="#0f172a", font=font(20))
    return img


def make_text_pdf(path: Path, pages: int = 3) -> None:
    doc = pymupdf.open()
    for page_index in range(pages):
        page = doc.new_page(width=595, height=842)
        page.insert_text((60, 80), f"DOCUPRO — Rapport annuel {2026}", fontsize=20, fontname="hebo")
        page.insert_text((60, 110), "Document de démonstration bilingue: français / English.", fontsize=10)
        page.insert_text((60, 150), f"Page {page_index + 1} of {pages}", fontsize=11)
        body = (
            "Ce document contient des accents (é, à, ç, œ), des tirets, des guillemets « français » "
            "et des symboles mathématiques: x² + y² = z², ½, ¾, ∑, √, ≤, ≥, π ≈ 3,14159."
        )
        page.insert_textbox(pymupdf.Rect(60, 170, 535, 260), body, fontsize=10, align=0)
        # simple ruled table
        top = 300
        rows = [
            ["Produit", "Quantité", "Prix (€)", "Total (€)"],
            ["Stylo", "12", "1,50", "18,00"],
            ["Cahier A4", "8", "3,20", "25,60"],
            ["Classeur", "4", "7,90", "31,60"],
        ]
        col_widths = [150, 100, 110, 110]
        y = top
        for row in rows:
            x = 60
            for cell, width in zip(row, col_widths):
                page.draw_rect(pymupdf.Rect(x, y, x + width, y + 24), color=(0.7, 0.75, 0.8), width=0.6)
                page.insert_text((x + 6, y + 16), cell, fontsize=9)
                x += width
            y += 24
        if page_index == 0:
            page.insert_image(pymupdf.Rect(60, 460, 535, 720), stream=_png_bytes(chart_image()))
    doc.save(path)
    doc.close()


def make_scanned_pdf(path: Path, text: str = "Facture 2026\nTotal: 1 250,00 €\nClient: Élodie Bénin") -> None:
    img = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(img)
    y = 140
    for line in text.split("\n"):
        draw.text((120, y), line, fill="black", font=font(44))
        y += 90
    draw.rectangle([100, 100, 1140, 1740], outline="#94a3b8", width=3)
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(pymupdf.Rect(0, 0, 595, 842), stream=buffer.getvalue())
    doc.save(path)
    doc.close()


def make_mixed_pdf(path: Path) -> None:
    text_pdf = path.with_suffix(".tmp.pdf")
    make_text_pdf(text_pdf, pages=1)
    scanned = path.with_suffix(".tmp-scan.pdf")
    make_scanned_pdf(scanned)
    out = pymupdf.open()
    with pymupdf.open(text_pdf) as a:
        out.insert_pdf(a)
    with pymupdf.open(scanned) as b:
        out.insert_pdf(b)
    out.save(path)
    out.close()
    text_pdf.unlink(missing_ok=True)
    scanned.unlink(missing_ok=True)


def make_password_pdf(path: Path, password: str = "motdepasse") -> None:
    temp = path.with_suffix(".tmp.pdf")
    make_text_pdf(temp, pages=1)
    with pymupdf.open(temp) as doc:
        doc.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw=password, user_pw=password)
    temp.unlink(missing_ok=True)


def make_docx(path: Path, pages: int = 2) -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Inches, Pt

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    header = document.sections[0].header
    header.paragraphs[0].text = "DOCUPRO — Démonstration"
    footer = document.sections[0].footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.text = "Document confidentiel — page "
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    field_run = footer_paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    field_run._r.append(fld_begin)
    field_run._r.append(instr)
    field_run._r.append(fld_end)

    document.add_heading("Rapport de démonstration", level=1)
    document.add_paragraph(
        "Ce document de test contient des accents français (é, è, ç, œ), des tableaux, "
        "une image et des sauts de page."
    )
    document.add_heading("Résultats financiers", level=2)
    table = document.add_table(rows=4, cols=3)
    table.style = "Light Grid Accent 1"
    data = [
        ("Trimestre", "Chiffre d'affaires", "Marge"),
        ("T1", "12 500 €", "9 300 €"),
        ("T2", "15 200 €", "11 100 €"),
        ("T3", "13 800 €", "10 050 €"),
    ]
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.add_picture(io.BytesIO(_png_bytes(chart_image(480, 260))), width=Inches(4.2))
    for _ in range(pages - 1):
        document.add_page_break()
        document.add_heading("Annexe", level=2)
        document.add_paragraph("Contenu additionnel avec des caractères spéciaux: ² ½ √ ≤ π.")
    document.save(path)


def _png_bytes(img: Image.Image) -> bytes:
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    return buffer.getvalue()


def make_images(directory: Path) -> None:
    base = Image.new("RGB", (1000, 700), "white")
    draw = ImageDraw.Draw(base)
    draw.text((60, 80), "DOCUPRO OCR TEST", fill="black", font=font(56))
    draw.text((60, 200), "Rapport financier 2026", fill="black", font=font(40))
    draw.text((60, 300), "Total: 12 500 EUR", fill="black", font=font(40))
    draw.text((60, 400), "Client: Amelie Dubois", fill="black", font=font(40))
    base.save(directory / "ocr-sample.png")
    base.save(directory / "photo.jpg", quality=92)
    base.convert("RGB").resize((800, 560)).save(directory / "photo.webp", quality=88)
    transparent = Image.new("RGBA", (600, 400), (0, 0, 0, 0))
    draw2 = ImageDraw.Draw(transparent)
    draw2.rectangle([50, 50, 550, 350], fill=(37, 99, 235, 255))
    transparent.save(directory / "logo-alpha.png")
    try:
        from pillow_heif import register_heif_opener

        register_heif_opener()
        base.save(directory / "photo.heic", quality=85)
    except Exception:
        pass
    wide = Image.new("RGB", (4000, 3000), "white")
    ImageDraw.Draw(wide).text((200, 200), "Grande image", fill="black", font=font(120))
    wide.save(directory / "large.jpg", quality=80)


def make_sheets(directory: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Ventes"
    ws.append(["Mois", "Ventes", "Marge"])
    for month, sales, margin in [("Janvier", 12500, 9300), ("Février", 15200, 11100), ("Mars", 13800, 10050)]:
        ws.append([month, sales, margin])
    ws2 = wb.create_sheet("Notes")
    ws2.append(["Commentaire"])
    ws2.append(["Données de démonstration — accents é à ç"])
    wb.save(directory / "workbook.xlsx")

    (directory / "data.csv").write_text(
        "Nom;Ville;Montant\nÉlodie Bénin;Abomey-Calavi;1250,50\nJean Dupont;Paris;980,00\n"
        "Aïcha Diallo;Dakar;1400,75\n",
        encoding="utf-8",
    )
    (directory / "plain.txt").write_text(
        "DOCUPRO — fichier texte de démonstration.\n\n"
        "Ce texte contient des accents et des symboles: ² ½ √ ≤ π.\n" * 5,
        encoding="utf-8",
    )
    (directory / "page.html").write_text(
        """<!doctype html><html><head><meta charset="utf-8"><title>Test</title>
        <script>alert('should be stripped')</script></head>
        <body><h1>Facture</h1><p>Merci pour votre achat.</p>
        <table><tr><th>Article</th><th>Prix</th></tr><tr><td>Stylo</td><td>1,50 €</td></tr></table>
        <img src="https://example.com/tracker.png" onerror="alert(1)">
        </body></html>""",
        encoding="utf-8",
    )


def make_invalid(directory: Path) -> None:
    (directory / "corrupted.pdf").write_bytes(b"%PDF-1.7\n1 0 obj\n<< broken >>")
    (directory / "fake.pdf").write_bytes(_png_bytes(Image.new("RGB", (50, 50), "red")))
    (directory / "notes.pdf").write_text("This is plain text pretending to be a PDF.", encoding="utf-8")
    (directory / "empty.pdf").write_bytes(b"")
    (directory / "huge.exe").write_bytes(b"MZ\x90\x00" + b"\x00" * 2048)


def main() -> int:
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "tests/corpus/generated").resolve()
    out.mkdir(parents=True, exist_ok=True)
    make_text_pdf(out / "report-fr.pdf", pages=3)
    make_scanned_pdf(out / "scan-fr.pdf")
    make_mixed_pdf(out / "mixed.pdf")
    make_password_pdf(out / "protected.pdf")
    make_docx(out / "report.docx", pages=3)
    make_docx(out / "short.docx", pages=1)
    make_images(out)
    make_sheets(out)
    make_invalid(out)
    print(f"Corpus written to {out}")
    for path in sorted(out.iterdir()):
        print(f"  {path.name:24} {path.stat().st_size:>9} bytes")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
