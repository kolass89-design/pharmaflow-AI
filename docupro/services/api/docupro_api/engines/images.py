"""Image conversions, including OCR-based extraction to text and Word."""

from __future__ import annotations

import io
import logging
from pathlib import Path

from PIL import Image

from ..errors import ConversionFailed, ValidationError
from . import ocr
from . import validation as v
from .common import (
    build_pdf_from_images,
    clamp,
    flatten_to_rgb,
    jpeg_bytes,
    load_image,
    make_zip,
    png_bytes,
    unique_name,
)

log = logging.getLogger(__name__)

try:  # optional HEIC/HEIF support
    from pillow_heif import register_heif_opener

    register_heif_opener()
    HEIF_SUPPORTED = True
except Exception:  # pragma: no cover - optional dependency
    HEIF_SUPPORTED = False
    log.info("HEIF support not available (pillow-heif not installed)")

IMAGE_MIME = {
    "jpeg": "image/jpeg",
    "jpg": "image/jpeg",
    "png": "image/png",
    "webp": "image/webp",
    "gif": "image/gif",
    "bmp": "image/bmp",
    "tiff": "image/tiff",
}


def _extension_for(fmt: str) -> str:
    return "jpg" if fmt in {"jpeg", "jpg"} else fmt.lower()


def _save_image(img: Image.Image, destination: Path, fmt: str, quality: int) -> None:
    fmt = fmt.lower()
    if fmt in {"jpeg", "jpg"}:
        destination.write_bytes(jpeg_bytes(img, quality))
    elif fmt == "png":
        destination.write_bytes(png_bytes(img))
    elif fmt == "webp":
        buffer = io.BytesIO()
        flatten_to_rgb(img).save(buffer, format="WEBP", quality=int(clamp(quality, 30, 100)), method=4, optimize=True)
        destination.write_bytes(buffer.getvalue())
    else:
        raise ValidationError("Unsupported output image format.", code="invalid_option")


def _parse_order(spec: str | None, count: int) -> list[int]:
    if not spec or not str(spec).strip():
        return list(range(count))
    try:
        order = [int(part) - 1 for part in str(spec).replace(";", ",").split(",") if part.strip()]
    except ValueError as exc:
        raise ValidationError("Image order must look like 3,1,2.", code="invalid_order") from exc
    if sorted(order) != list(range(count)):
        raise ValidationError(
            f"Image order must list each of the {count} files exactly once.",
            code="invalid_order",
        )
    return order


def _maybe_zip(ctx, outputs: list[Path], names: list[str], archive_name: str, meta: dict) -> list:
    if len(outputs) == 1:
        return [
            ctx.emit(
                outputs[0],
                names[0],
                IMAGE_MIME.get(outputs[0].suffix.lstrip(".").lower(), "application/octet-stream"),
                meta=meta,
            )
        ]
    if not ctx.opt_bool("zip", True):
        return [
            ctx.emit(p, n, IMAGE_MIME.get(p.suffix.lstrip(".").lower(), "application/octet-stream"), meta=meta)
            for p, n in zip(outputs, names)
        ]
    archive = ctx.outdir / "images.zip"
    make_zip(outputs, archive, names)
    v.validate_zip(archive)
    return [ctx.emit(archive, archive_name, "application/zip", meta={**meta, "files": len(outputs)})]


# --- image -> PDF -----------------------------------------------------------


def images_to_pdf(ctx) -> list:
    order = _parse_order(ctx.opt("order", ""), len(ctx.files))
    page_size = str(ctx.opt("page_size", "a4"))
    orientation = str(ctx.opt("orientation", "auto"))
    margin_mm = clamp(ctx.opt_float("margin", 10.0), 0.0, 50.0)
    fit_mode = str(ctx.opt("fit_mode", "contain"))
    quality = ctx.opt_int("quality", 88, 40, 100)
    if page_size not in {"fit", "a4", "letter", "legal", "a3", "a5"}:
        raise ValidationError("Unsupported page size.", code="invalid_option")
    if orientation not in {"auto", "portrait", "landscape"}:
        raise ValidationError("Unsupported orientation.", code="invalid_option")
    if fit_mode not in {"contain", "cover", "stretch"}:
        raise ValidationError("Unsupported fit mode.", code="invalid_option")

    images: list[Image.Image] = []
    for position, index in enumerate(order):
        ctx.check_cancel()
        source = ctx.files[index]
        img = load_image(source.path, ctx.settings.max_pixels)
        images.append(img)
        ctx.report(0.05 + 0.5 * (position + 1) / len(order), f"Reading image {position + 1}")
    if not images:
        raise ConversionFailed("No images could be read.")

    doc = build_pdf_from_images(images, page_size, orientation, margin_mm, fit_mode, quality)
    target = ctx.outdir / "images.pdf"
    ctx.report(0.75, "Building PDF")
    try:
        doc.save(target, garbage=3, deflate=True, deflate_images=True)
    finally:
        pages = doc.page_count
        doc.close()
    v.validate_pdf(target, min_pages=len(images))
    base = ctx.files[0].stem if len(ctx.files) == 1 else "images"
    return [
        ctx.emit(
            target,
            f"{base}.pdf",
            "application/pdf",
            meta={
                "pages": pages,
                "images": len(images),
                "pageSize": page_size,
                "orientation": orientation,
                "fitMode": fit_mode,
                "marginMm": margin_mm,
            },
        )
    ]


# --- image format conversion ------------------------------------------------


def image_convert(ctx) -> list:
    target_format = str(ctx.opt("target_format", "png")).lower()
    if target_format not in {"jpeg", "png", "webp"}:
        raise ValidationError("Choose JPEG, PNG or WebP as the output format.", code="invalid_option")
    quality = ctx.opt_int("quality", 90, 40, 100)
    outputs: list[Path] = []
    names: list[str] = []
    used: set[str] = set()
    for index, source in enumerate(ctx.files):
        ctx.check_cancel()
        img = load_image(source.path, ctx.settings.max_pixels)
        ext = _extension_for(target_format)
        name = unique_name(f"{source.stem}", used, ext)
        target = ctx.outdir / f"converted-{index + 1:03d}.{ext}"
        _save_image(img, target, target_format, quality)
        v.validate_image(target, expected_format=target_format)
        outputs.append(target)
        names.append(name)
        ctx.report(0.05 + 0.9 * (index + 1) / len(ctx.files), f"Converted {index + 1} of {len(ctx.files)}")
    meta = {"format": target_format, "quality": quality, "count": len(outputs)}
    archive_name = f"{ctx.files[0].stem}.{_extension_for(target_format)}"
    if len(outputs) > 1:
        archive_name = f"images-{_extension_for(target_format)}.zip"
    return _maybe_zip(ctx, outputs, names, archive_name, meta)


# --- resize -----------------------------------------------------------------


def image_resize(ctx) -> list:
    mode = str(ctx.opt("mode", "max_edge"))
    if mode not in {"percent", "width", "max_edge", "exact"}:
        raise ValidationError("Unsupported resize mode.", code="invalid_option")
    percent = clamp(ctx.opt_float("percent", 50), 1, 400)
    width = ctx.opt_int("width", 1920, 16, 20000)
    height = ctx.opt_int("height", 1080, 16, 20000)
    keep_aspect = ctx.opt_bool("keep_aspect", True)
    target_format = str(ctx.opt("target_format", "keep")).lower()
    quality = ctx.opt_int("quality", 92, 40, 100)
    outputs: list[Path] = []
    names: list[str] = []
    used: set[str] = set()
    for index, source in enumerate(ctx.files):
        ctx.check_cancel()
        img = load_image(source.path, ctx.settings.max_pixels)
        src_fmt = (img.format or source.path.suffix.lstrip(".") or "png").lower()
        fmt = target_format if target_format != "keep" else ("jpeg" if src_fmt in {"jpeg", "jpg"} else src_fmt)
        if fmt not in {"jpeg", "png", "webp"}:
            fmt = "png"
        if mode == "percent":
            size = (max(1, int(img.width * percent / 100)), max(1, int(img.height * percent / 100)))
        elif mode == "width":
            ratio = width / max(1, img.width)
            size = (width, max(1, int(img.height * ratio))) if keep_aspect else (width, height)
        elif mode == "exact":
            size = (width, height)
        else:  # max_edge
            scale = min(1.0, width / max(1, max(img.width, img.height)))
            size = (max(1, int(img.width * scale)), max(1, int(img.height * scale)))
        if size[0] * size[1] > ctx.settings.max_pixels:
            raise ValidationError("The requested size exceeds the allowed pixel budget.", code="image_too_large")
        resized = img.resize(size, Image.LANCZOS)
        ext = _extension_for(fmt)
        name = unique_name(f"{source.stem}-{size[0]}x{size[1]}", used, ext)
        target = ctx.outdir / f"resized-{index + 1:03d}.{ext}"
        _save_image(resized, target, fmt, quality)
        v.validate_image(target)
        outputs.append(target)
        names.append(name)
        ctx.report(0.05 + 0.9 * (index + 1) / len(ctx.files), f"Resized {index + 1} of {len(ctx.files)}")
    meta = {"mode": mode, "count": len(outputs)}
    archive_name = names[0] if len(names) == 1 else "resized-images.zip"
    return _maybe_zip(ctx, outputs, names, archive_name, meta)


# --- compression ------------------------------------------------------------


def image_compress(ctx) -> list:
    quality = ctx.opt_int("quality", 70, 20, 95)
    target_format = str(ctx.opt("target_format", "jpeg")).lower()
    if target_format not in {"jpeg", "webp"}:
        target_format = "jpeg"
    max_edge = ctx.opt_int("max_edge", 0, 0, 12000)
    outputs: list[Path] = []
    names: list[str] = []
    used: set[str] = set()
    original_total = 0
    new_total = 0
    for index, source in enumerate(ctx.files):
        ctx.check_cancel()
        original_total += source.size
        img = load_image(source.path, ctx.settings.max_pixels)
        if max_edge and max(img.size) > max_edge:
            scale = max_edge / max(img.size)
            img = img.resize((max(1, int(img.width * scale)), max(1, int(img.height * scale))), Image.LANCZOS)
        ext = _extension_for(target_format)
        name = unique_name(f"{source.stem}-compressed", used, ext)
        target = ctx.outdir / f"compressed-{index + 1:03d}.{ext}"
        _save_image(img, target, target_format, quality)
        v.validate_image(target)
        new_total += target.stat().st_size
        outputs.append(target)
        names.append(name)
        ctx.report(0.05 + 0.9 * (index + 1) / len(ctx.files), f"Compressed {index + 1} of {len(ctx.files)}")
    meta = {
        "quality": quality,
        "format": target_format,
        "originalSize": original_total,
        "compressedSize": new_total,
        "savedPercent": round(max(0.0, 1 - new_total / original_total) * 100, 1) if original_total else 0,
    }
    archive_name = names[0] if len(names) == 1 else "compressed-images.zip"
    return _maybe_zip(ctx, outputs, names, archive_name, meta)


# --- rotate -----------------------------------------------------------------


def image_rotate(ctx) -> list:
    angle = ctx.opt_int("angle", 90)
    if angle not in {90, 180, 270}:
        raise ValidationError("Rotation must be 90, 180 or 270 degrees.", code="invalid_option")
    flip = ctx.opt_bool("flip", False)
    outputs: list[Path] = []
    names: list[str] = []
    used: set[str] = set()
    for index, source in enumerate(ctx.files):
        ctx.check_cancel()
        img = load_image(source.path, ctx.settings.max_pixels)
        rotated = img.rotate(-angle, expand=True)
        if flip:
            from PIL import ImageOps

            rotated = ImageOps.mirror(rotated)
        src_fmt = (img.format or "PNG").lower()
        fmt = "jpeg" if src_fmt in {"jpeg", "jpg"} else ("webp" if src_fmt == "webp" else "png")
        ext = _extension_for(fmt)
        name = unique_name(f"{source.stem}-rotated", used, ext)
        target = ctx.outdir / f"rotated-{index + 1:03d}.{ext}"
        _save_image(rotated, target, fmt, 92)
        v.validate_image(target)
        outputs.append(target)
        names.append(name)
        ctx.report(0.05 + 0.9 * (index + 1) / len(ctx.files), f"Rotated {index + 1} of {len(ctx.files)}")
    meta = {"angle": angle, "flip": flip}
    archive_name = names[0] if len(names) == 1 else "rotated-images.zip"
    return _maybe_zip(ctx, outputs, names, archive_name, meta)


# --- OCR: image -> text / Word ---------------------------------------------


def image_to_text(ctx) -> list:
    language = str(ctx.opt("ocr_language", "auto"))
    ocr.require_ocr_engine()
    total_confidence: list[float] = []
    chunks: list[str] = []
    for index, source in enumerate(ctx.files):
        ctx.check_cancel()
        img = load_image(source.path, ctx.settings.max_pixels)
        lines = ocr.ocr_lines(img, language)
        total_confidence.extend(line.confidence for line in lines)
        paragraphs = ocr.lines_to_paragraphs(lines)
        text = "\n\n".join(paragraphs).strip()
        if len(ctx.files) > 1:
            chunks.append(f"--- {source.original_name} ---\n\n{text}")
        elif text:
            chunks.append(text)
        ctx.report(0.05 + 0.9 * (index + 1) / len(ctx.files), f"Recognised {index + 1} of {len(ctx.files)}")
    content = "\n\n".join(chunks).strip()
    if not content:
        raise ConversionFailed(
            "No text was detected in the image. Try a sharper scan or a higher-resolution photo.",
            code="no_text",
        )
    target = ctx.outdir / "text.txt"
    target.write_text(content + "\n", encoding="utf-8")
    v.validate_text(target)
    confidence = round(sum(total_confidence) / len(total_confidence), 3) if total_confidence else None
    return [
        ctx.emit(
            target,
            f"{ctx.files[0].stem}.txt" if len(ctx.files) == 1 else "image-text.txt",
            "text/plain; charset=utf-8",
            meta={"characters": len(content), "confidence": confidence, "language": language},
        )
    ]


def image_to_word(ctx) -> list:
    from docx import Document
    from docx.shared import Pt, RGBColor

    language = str(ctx.opt("ocr_language", "auto"))
    include_image = ctx.opt_bool("include_image", False)
    ocr.require_ocr_engine()
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    confidences: list[float] = []
    wrote_any = False
    for index, source in enumerate(ctx.files):
        ctx.check_cancel()
        img = load_image(source.path, ctx.settings.max_pixels)
        lines = ocr.ocr_lines(img, language)
        confidences.extend(line.confidence for line in lines)
        paragraphs = ocr.lines_to_paragraphs(lines)
        if len(ctx.files) > 1:
            document.add_heading(source.stem, level=2)
        for paragraph in paragraphs:
            if paragraph.strip():
                document.add_paragraph(paragraph.strip())
                wrote_any = True
        if include_image:
            document.add_picture(str(source.path), width=None) if len(ctx.files) == 1 else None
        if index < len(ctx.files) - 1:
            document.add_page_break()
        ctx.report(0.05 + 0.85 * (index + 1) / len(ctx.files), f"Recognised {index + 1} of {len(ctx.files)}")
    if not wrote_any:
        raise ConversionFailed(
            "No text was detected in the image, so an editable document could not be created.",
            code="no_text",
        )
    note = document.add_paragraph()
    run = note.add_run(
        "Editable text was reconstructed from an image with OCR. Layout, fonts and exact spacing "
        "cannot be reproduced from a scanned image."
    )
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    target = ctx.outdir / "image.docx"
    document.save(str(target))
    items = v.validate_docx(target)
    confidence = round(sum(confidences) / len(confidences), 3) if confidences else None
    return [
        ctx.emit(
            target,
            f"{ctx.files[0].stem}.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            meta={"blocks": items, "confidence": confidence, "language": language, "images": len(ctx.files)},
        )
    ]
