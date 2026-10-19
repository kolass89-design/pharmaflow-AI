"""Conversions that start from a PDF."""

from __future__ import annotations

import logging
import re
from pathlib import Path

import pymupdf

from ..errors import ConversionFailed
from . import ocr
from . import validation as v
from .common import (
    make_zip,
    normalize_ligatures,
    open_pdf,
    parse_page_selection,
    render_page_bytes,
    require_page_budget,
)

log = logging.getLogger(__name__)
PDF_MIME = "application/pdf"


# --- PDF -> Word ------------------------------------------------------------


def _document_text_ratio(doc: pymupdf.Document) -> float:
    """Share of pages that carry real, extractable text."""
    if doc.page_count == 0:
        return 0.0
    textual = 0
    for page in doc:
        if len(page.get_text("text").strip()) >= 30:
            textual += 1
    return textual / doc.page_count


def _ocr_docx(ctx, doc: pymupdf.Document, language: str, note: str) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    header = document.add_paragraph()
    run = header.add_run(note)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT

    total = doc.page_count
    for index, page in enumerate(doc):
        ctx.check_cancel()
        ctx.report(0.1 + 0.75 * index / max(1, total), f"Reading page {index + 1} of {total}")
        pix = page.get_pixmap(dpi=210, colorspace=pymupdf.csRGB, alpha=False)
        import io

        from PIL import Image

        image = Image.open(io.BytesIO(pix.tobytes("png")))
        lines = ocr.ocr_lines(image, language)
        paragraphs = ocr.lines_to_paragraphs(lines)
        # Preserve simple indentation based on the left edge of each paragraph.
        page_width = max(1.0, page.rect.width)
        first_line_left: dict[str, float] = {}
        for line in lines:
            text = line.text.strip()
            if text and text not in first_line_left:
                first_line_left[text] = line.bbox[0]
        for paragraph in paragraphs:
            p = document.add_paragraph(paragraph)
            left = first_line_left.get(paragraph.split("\n")[0][: len(paragraph)])
            if left is not None and left / page_width > 0.12:
                p.paragraph_format.left_indent = Pt(min(72.0, left / page_width * 300))
        if index < total - 1:
            document.add_page_break()
    target = ctx.outdir / "converted.docx"
    document.save(str(target))
    return target


def _normalize_docx_text(path: Path) -> None:
    """Replace presentation-form ligatures left behind by the layout analyser.

    pdf2docx copies glyphs such as U+FB01 straight out of the PDF, which makes
    words like "fichier" unsearchable in the resulting Word document.
    """
    import docx

    document = docx.Document(str(path))
    changed = False

    def fix(paragraphs) -> None:
        nonlocal changed
        for paragraph in paragraphs:
            for run in paragraph.runs:
                cleaned = normalize_ligatures(run.text)
                if cleaned != run.text:
                    run.text = cleaned
                    changed = True

    fix(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                fix(cell.paragraphs)
    for section in document.sections:
        for part in (section.header, section.footer):
            fix(part.paragraphs)
    if changed:
        document.save(str(path))


def pdf_to_word(ctx) -> list:
    source = ctx.files[0]
    mode = str(ctx.opt("mode", "auto"))
    language = str(ctx.opt("ocr_language", "auto"))
    with open_pdf(source.path) as doc:
        require_page_budget(doc.page_count, ctx.settings)
        ratio = _document_text_ratio(doc)
        page_count = doc.page_count
        use_ocr = mode == "ocr" or (mode == "auto" and ratio < 0.35)

        if use_ocr:
            ocr.require_ocr_engine()
            target = _ocr_docx(
                ctx,
                doc,
                language,
                "Reconstructed from a scanned PDF with OCR. Text is editable, but exact "
                "layout, tables and fonts may differ from the original document.",
            )
            method = "ocr"
        else:
            ctx.report(0.15, "Analysing layout")
            from pdf2docx import Converter

            target = ctx.outdir / "converted.docx"
            try:
                converter = Converter(str(source.path))
                try:
                    converter.convert(str(target))
                finally:
                    converter.close()
            except Exception as exc:
                log.warning("Layout conversion failed, falling back to text extraction: %s", exc)
                if not ctx.opt_bool("text_fallback", True):
                    raise ConversionFailed(
                        "This PDF's layout could not be converted to Word. Try the OCR mode."
                    ) from exc
                target = _ocr_docx(
                    ctx,
                    doc,
                    language,
                    "Reconstructed from a PDF whose layout could not be analysed automatically. "
                    "Paragraph text is preserved; complex layout may differ.",
                )
                method = "ocr-fallback"
            else:
                method = "layout"

    if method == "layout":
        _normalize_docx_text(target)
    items = v.validate_docx(target)
    ctx.report(0.98, "Validating output")
    return [
        ctx.emit(
            target,
            f"{source.stem}.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            meta={"pages": page_count, "method": method, "blocks": items},
        )
    ]


# --- PDF -> images ----------------------------------------------------------


def _render_pages(ctx, doc, fmt: str, dpi: int, selection: list[int], quality: int):
    outputs: list[Path] = []
    names: list[str] = []
    stem = ctx.files[0].stem
    ext = "jpg" if fmt == "jpeg" else "png"
    for position, index in enumerate(selection):
        ctx.check_cancel()
        page = doc.load_page(index)
        data, width, height = render_page_bytes(page, dpi, fmt, ctx.settings.max_pixels)
        target = ctx.outdir / f"page-{index + 1:04d}.{ext}"
        target.write_bytes(data)
        v.validate_image(target, expected_format=fmt)
        outputs.append(target)
        names.append(f"{stem}-page-{index + 1:03d}.{ext}")
        ctx.report(0.05 + 0.9 * (position + 1) / len(selection), f"Rendering page {index + 1}")
    return outputs, names, width, height


def pdf_to_images(ctx, fmt: str) -> list:
    source = ctx.files[0]
    dpi = ctx.opt_int("dpi", 150, 50, 600)
    quality = ctx.opt_int("quality", 88, 30, 100)
    with open_pdf(source.path) as doc:
        require_page_budget(doc.page_count, ctx.settings)
        selection = parse_page_selection(ctx.opt("pages", ""), doc.page_count)
        outputs, names, width, height = _render_pages(ctx, doc, fmt, dpi, selection, quality)
    ext = "jpg" if fmt == "jpeg" else "png"
    media = "image/jpeg" if fmt == "jpeg" else "image/png"
    meta = {"pages": len(outputs), "dpi": dpi, "width": width, "height": height}

    if len(outputs) == 1:
        return [ctx.emit(outputs[0], names[0], media, meta=meta)]
    if not ctx.opt_bool("zip", True):
        return [
            ctx.emit(path, name, media, meta={"page": i + 1, "dpi": dpi})
            for i, (path, name) in enumerate(zip(outputs, names))
        ]
    archive = ctx.outdir / "pages.zip"
    make_zip(outputs, archive, names)
    v.validate_zip(archive)
    return [
        ctx.emit(
            archive,
            f"{source.stem}-{ext}-{dpi}dpi.zip",
            "application/zip",
            meta={**meta, "files": len(outputs)},
        )
    ]


def pdf_to_jpeg(ctx) -> list:
    return pdf_to_images(ctx, "jpeg")


def pdf_to_png(ctx) -> list:
    return pdf_to_images(ctx, "png")


def pdf_extract_images(ctx) -> list:
    """Extract embedded raster images (not page renderings)."""
    source = ctx.files[0]
    outputs: list[Path] = []
    names: list[str] = []
    with open_pdf(source.path) as doc:
        require_page_budget(doc.page_count, ctx.settings)
        seen: set[int] = set()
        total_pages = doc.page_count
        for pno in range(total_pages):
            ctx.check_cancel()
            for info in doc.get_page_images(pno, full=True):
                xref = info[0]
                if xref in seen:
                    continue
                seen.add(xref)
                try:
                    payload = doc.extract_image(xref)
                except Exception:
                    continue
                if not payload or not payload.get("image"):
                    continue
                ext = (payload.get("ext") or "png").lower()
                if ext == "jpx":
                    ext = "jp2"
                if ext not in {"png", "jpg", "jpeg", "jp2", "webp", "tiff", "bmp", "gif"}:
                    ext = "png"
                target = ctx.outdir / f"image-{len(outputs) + 1:04d}.{ext}"
                target.write_bytes(payload["image"])
                outputs.append(target)
                names.append(f"{source.stem}-image-{len(outputs):03d}.{ext}")
            ctx.report(0.1 + 0.8 * (pno + 1) / doc.page_count, f"Scanning page {pno + 1}")
    if not outputs:
        raise ConversionFailed("This PDF contains no embedded images. Use PDF to PNG to render the pages instead.")
    archive = ctx.outdir / "images.zip"
    make_zip(outputs, archive, names)
    v.validate_zip(archive)
    return [
        ctx.emit(
            archive,
            f"{source.stem}-images.zip",
            "application/zip",
            meta={"files": len(outputs), "pages": total_pages},
        )
    ]


# --- PDF -> text ------------------------------------------------------------


def pdf_to_text(ctx) -> list:
    source = ctx.files[0]
    with open_pdf(source.path) as doc:
        require_page_budget(doc.page_count, ctx.settings)
        selection = parse_page_selection(ctx.opt("pages", ""), doc.page_count)
        markers = ctx.opt_bool("page_markers", True)
        ocr_fallback = ctx.opt_bool("ocr_fallback", False)
        language = str(ctx.opt("ocr_language", "auto"))
        chunks: list[str] = []
        for position, index in enumerate(selection):
            ctx.check_cancel()
            page = doc.load_page(index)
            text = normalize_ligatures(page.get_text("text", sort=True).strip())
            if not text and ocr_fallback:
                ocr.require_ocr_engine()
                pix = page.get_pixmap(dpi=220, colorspace=pymupdf.csRGB, alpha=False)
                import io

                from PIL import Image

                image = Image.open(io.BytesIO(pix.tobytes("png")))
                text = ocr.ocr_to_text(image, language)
            if markers:
                chunks.append(f"--- Page {index + 1} ---\n\n{text}".rstrip())
            else:
                chunks.append(text)
            ctx.report(0.05 + 0.85 * (position + 1) / len(selection), f"Reading page {index + 1}")
        content = "\n\n".join(c for c in chunks if c).strip()
    if not content:
        raise ConversionFailed(
            "No extractable text was found. This looks like a scanned PDF: try PDF to Word with OCR, "
            "or enable OCR fallback.",
            code="no_text",
        )
    target = ctx.outdir / "content.txt"
    target.write_text(content + "\n", encoding="utf-8")
    size = v.validate_text(target)
    return [
        ctx.emit(
            target,
            f"{source.stem}.txt",
            "text/plain; charset=utf-8",
            meta={"characters": size, "pages": len(selection)},
        )
    ]


# --- PDF -> searchable PDF --------------------------------------------------


def _invisible_text(page, lines, scale: float) -> int:
    """Stamp an invisible, selectable text layer from OCR lines.

    Returns the number of lines actually embedded. ``insert_textbox`` returns a
    negative value (without raising) when the box is too small, so the return
    code is checked before counting a line as written.
    """
    from .fonts import page_font_alias

    alias, font = page_font_alias(page)
    written = 0
    for line in lines:
        text = re.sub(r"\s+", " ", line.text).strip()
        if not text:
            continue
        if alias != "helv":
            # Drop characters the embedded font cannot represent instead of
            # silently writing replacement glyphs into the text layer.
            text = "".join(ch for ch in text if ch.isspace() or font.has_glyph(ord(ch)))
            text = text.strip()
            if not text:
                continue
        x0, y0, x1, y1 = (value / scale for value in line.bbox)
        width = max(1.0, x1 - x0)
        height = max(1.0, y1 - y0)
        fontsize = max(3.0, height * 0.8)
        measured = font.text_length(text, fontsize=fontsize)
        if measured > width > 0:
            fontsize = max(2.5, fontsize * (width / measured))
        placed = False
        for pad in (2.0, 6.0, 12.0):
            rect = pymupdf.Rect(x0, y0 - pad / 2, x1 + pad, y1 + pad)
            try:
                if (
                    page.insert_textbox(
                        rect,
                        text,
                        fontsize=fontsize,
                        fontname=alias,
                        color=(0, 0, 0),
                        render_mode=3,
                        overlay=True,
                    )
                    >= 0
                ):
                    placed = True
                    break
            except Exception:
                continue
        if not placed:
            try:
                page.insert_text(
                    pymupdf.Point(x0, y1 - 0.5),
                    text,
                    fontsize=fontsize,
                    fontname=alias,
                    render_mode=3,
                    overlay=True,
                )
                placed = True
            except Exception:  # pragma: no cover
                placed = False
        if placed:
            written += 1
    return written


def pdf_to_searchable(ctx) -> list:
    source = ctx.files[0]
    language = str(ctx.opt("ocr_language", "auto"))
    dpi = ctx.opt_int("dpi", 300, 120, 400)
    force = ctx.opt_bool("force", False)
    ocr.require_ocr_engine()
    with open_pdf(source.path) as doc:
        require_page_budget(doc.page_count, ctx.settings)
        processed = 0
        skipped = 0
        for index in range(doc.page_count):
            ctx.check_cancel()
            page = doc.load_page(index)
            has_text = len(page.get_text("text").strip()) >= 10
            if has_text and not force:
                skipped += 1
                ctx.report(0.05 + 0.9 * (index + 1) / doc.page_count, f"Skipped page {index + 1} (already searchable)")
                continue
            zoom = dpi / 72.0
            pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), colorspace=pymupdf.csRGB, alpha=False)
            import io

            from PIL import Image

            image = Image.open(io.BytesIO(pix.tobytes("png")))
            lines = ocr.ocr_lines(image, language)
            written = _invisible_text(page, lines, zoom)
            if written and not page.get_text("text").strip():  # pragma: no cover - defensive
                raise ConversionFailed(
                    "The OCR text layer could not be verified on this page.",
                    code="ocr_embed_failed",
                )
            if written == 0:
                raise ConversionFailed(
                    "OCR found text on a page but it could not be embedded in the PDF text layer.",
                    code="ocr_embed_failed",
                )
            processed += 1
            ctx.report(0.05 + 0.9 * (index + 1) / doc.page_count, f"Recognised page {index + 1}")
        target = ctx.outdir / "searchable.pdf"
        doc.set_metadata({**(doc.metadata or {}), "producer": "DOCUPRO", "creator": "DOCUPRO"})
        doc.save(target, garbage=4, deflate=True, deflate_images=True, clean=False)
        pages = doc.page_count
    v.validate_pdf(target)
    if processed == 0:
        raise ConversionFailed(
            "Every page already contains text; no OCR layer was required.",
            code="already_searchable",
            details={"pages": pages, "skipped": skipped},
        )
    return [
        ctx.emit(
            target,
            f"{source.stem}-searchable.pdf",
            PDF_MIME,
            meta={"pages": pages, "ocrPages": processed, "skipped": skipped, "dpi": dpi, "language": language},
        )
    ]


# --- PDF -> Excel -----------------------------------------------------------


def pdf_to_excel(ctx) -> list:
    source = ctx.files[0]
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    total_tables = 0
    with open_pdf(source.path) as doc:
        require_page_budget(doc.page_count, ctx.settings)
        for pno in range(doc.page_count):
            ctx.check_cancel()
            page = doc.load_page(pno)
            try:
                found = page.find_tables()
                tables = list(found.tables)
            except Exception:  # pragma: no cover - engine dependent
                tables = []
            for t_index, table in enumerate(tables):
                rows = table.extract()
                rows = [row for row in rows if any((cell or "").strip() for cell in row)] if rows else []
                if len(rows) < 2:
                    continue
                total_tables += 1
                title = f"Page {pno + 1}" + (f" T{t_index + 1}" if len(tables) > 1 else "")
                ws = wb.create_sheet(title=title[:31])
                for r_index, row in enumerate(rows, start=1):
                    for c_index, value in enumerate(row, start=1):
                        cell = ws.cell(row=r_index, column=c_index, value=(value or "").strip())
                        if r_index == 1:
                            cell.font = Font(bold=True)
                for column in ws.columns:
                    width = max((len(str(c.value or "")) for c in column), default=8)
                    ws.column_dimensions[column[0].column_letter].width = min(60, max(10, width + 2))
                ws.freeze_panes = "A2"
            ctx.report(0.1 + 0.8 * (pno + 1) / doc.page_count, f"Scanning page {pno + 1} for tables")
    if total_tables == 0:
        raise ConversionFailed(
            "No tables were detected in this PDF. PDF to Excel works best with ruled tables; "
            "try PDF to Word or PDF to text for other layouts.",
            code="no_tables",
        )
    target = ctx.outdir / "extracted.xlsx"
    wb.save(str(target))
    sheets = v.validate_xlsx(target)
    return [
        ctx.emit(
            target,
            f"{source.stem}-tables.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            meta={"tables": total_tables, "sheets": sheets},
        )
    ]


# --- PDF -> one PDF per page ------------------------------------------------


def pdf_to_single_pages(ctx) -> list:
    from .common import copy_pages

    source = ctx.files[0]
    outputs: list[Path] = []
    names: list[str] = []
    with open_pdf(source.path) as doc:
        require_page_budget(doc.page_count, ctx.settings)
        for index in range(doc.page_count):
            ctx.check_cancel()
            part = copy_pages(doc, [index])
            target = ctx.outdir / f"page-{index + 1:04d}.pdf"
            part.save(target, garbage=2, deflate=True)
            part.close()
            v.validate_pdf(target)
            outputs.append(target)
            names.append(f"{source.stem}-page-{index + 1:03d}.pdf")
            ctx.report(0.05 + 0.9 * (index + 1) / doc.page_count, f"Splitting page {index + 1}")
    if len(outputs) == 1:
        return [ctx.emit(outputs[0], names[0], PDF_MIME, meta={"pages": 1})]
    archive = ctx.outdir / "pages.zip"
    make_zip(outputs, archive, names)
    v.validate_zip(archive)
    return [
        ctx.emit(
            archive,
            f"{source.stem}-pages.zip",
            "application/zip",
            meta={"files": len(outputs)},
        )
    ]
