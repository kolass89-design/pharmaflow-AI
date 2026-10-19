"""Pure-Python DOCX renderer used when LibreOffice is unavailable.

It walks the document body in order, converts styles, runs, tables and images
into HTML, lays the HTML out with PyMuPDF's Story engine, then stamps the
document's header, footer and page numbers onto the finished pages. This keeps
text selectable and preserves the structural features users care about
(headings, tables, images, page breaks, page size, margins, numbering) even
without an office suite installed.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf

from ..errors import ConversionFailed
from . import validation as v

log = logging.getLogger(__name__)

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
EMU_PER_POINT = 12700.0


@dataclass
class RenderInfo:
    method: str = "python-html"
    pages: int = 0
    header: str = ""
    footer: str = ""
    has_page_field: bool = False
    page_width: float = 595.28
    page_height: float = 841.89
    margins: tuple[float, float, float, float] = (56.0, 56.0, 56.0, 56.0)
    warnings: list[str] = field(default_factory=list)


def _escape(text: str) -> str:
    return (text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _pt(value) -> float | None:
    return float(value) / EMU_PER_POINT if value is not None else None


def _color(run) -> str | None:
    try:
        rgb = run.font.color.rgb
        if rgb is not None:
            return f"#{rgb}"
    except Exception:  # pragma: no cover - theme colours
        return None
    return None


def _run_html(run, assets: dict[str, bytes]) -> str:
    """Convert a python-docx run to HTML, including images and line breaks."""
    pieces: list[str] = []
    element = run._element
    # Preserve line breaks and tabs inside the run.
    for child in element.iter():
        if child.tag == f"{W}br":
            if child.get(f"{W}type") == "page":
                pieces.append('<span class="page-break"></span>')
            else:
                pieces.append("<br/>")
        elif child.tag == f"{W}tab":
            pieces.append("&nbsp;&nbsp;&nbsp;&nbsp;")
    text = _escape(run.text)
    if not text:
        text = ""
    if run.bold:
        text = f"<strong>{text}</strong>" if text else text
    if run.italic:
        text = f"<em>{text}</em>" if text else text
    if run.underline:
        text = f"<u>{text}</u>" if text else text
    if run.font.strike:
        text = f"<s>{text}</s>" if text else text
    if run.font.superscript:
        text = f"<sup>{text}</sup>" if text else text
    if run.font.subscript:
        text = f"<sub>{text}</sub>" if text else text
    size = run.font.size.pt if run.font.size else None
    color = _color(run)
    style = ""
    if size:
        style += f"font-size:{size:.1f}pt;"
    if color:
        style += f"color:{color};"
    if run.font.highlight_color is not None:
        style += "background:#fef08a;"
    if style and text:
        text = f'<span style="{style}">{text}</span>'
    if text:
        pieces.append(text)

    for drawing in element.iter(f"{W}drawing"):
        blip = None
        for node in drawing.iter(f"{A}blip"):
            blip = node
            break
        if blip is None:
            continue
        rel_id = blip.get(f"{R}embed")
        if not rel_id:
            continue
        try:
            part = run.part.related_parts[rel_id]
        except (KeyError, AttributeError):
            continue
        blob = part.blob
        ext = (part.partname.ext or "png").lower()
        name = f"image-{len(assets) + 1:03d}.{ext if ext in {'png', 'jpg', 'jpeg', 'gif', 'bmp', 'tiff', 'webp'} else 'png'}"
        assets[name] = blob
        width = None
        for extent in drawing.iter(f"{WP}extent"):
            width = _pt(int(extent.get("cx"))) if extent.get("cx") else None
            break
        style_attr = f' width="{width:.0f}"' if width else ""
        pieces.append(f'<br/><img src="{name}"{style_attr}/><br/>')
    return "".join(pieces)


def _table_html(table, assets: dict[str, bytes]) -> str:
    rows_html: list[str] = []
    for row_index, row in enumerate(table.rows):
        cells_html: list[str] = []
        seen: set[int] = set()
        for cell in row.cells:
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            content_parts = []
            for paragraph in cell.paragraphs:
                html = "".join(_run_html(run, assets) for run in paragraph.runs)
                if html.strip():
                    content_parts.append(html)
            for nested in cell.tables:
                content_parts.append(_table_html(nested, assets))
            tag = "th" if row_index == 0 else "td"
            cells_html.append(f"<{tag}>{''.join(content_parts) or '&nbsp;'}</{tag}>")
        rows_html.append("<tr>" + "".join(cells_html) + "</tr>")
    return f"<table>{''.join(rows_html)}</table>"


def _iter_blocks(document):
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _style_kind(paragraph) -> tuple[str, str]:
    try:
        name = (paragraph.style.name or "").lower()
    except Exception:  # pragma: no cover
        name = ""
    if name.startswith("heading") or name.startswith("title"):
        level = name[-1] if name[-1].isdigit() else "1"
        if name.startswith("title"):
            return "h1", name
        return f"h{min(6, max(1, int(level)))}", name
    if "list bullet" in name or name.startswith("bullet"):
        return "li-bullet", name
    if "list number" in name or name.startswith("number"):
        return "li-number", name
    if "quote" in name or "intense quote" in name:
        return "blockquote", name
    return "p", name


def build_html(source: Path) -> tuple[str, dict[str, bytes], RenderInfo]:
    import docx

    try:
        document = docx.Document(str(source))
    except Exception as exc:
        raise ConversionFailed("This Word document could not be opened.") from exc

    info = RenderInfo()
    try:
        section = document.sections[0]
        info.page_width = _pt(section.page_width) or info.page_width
        info.page_height = _pt(section.page_height) or info.page_height
        info.margins = (
            _pt(section.top_margin) or 56.0,
            _pt(section.bottom_margin) or 56.0,
            _pt(section.left_margin) or 56.0,
            _pt(section.right_margin) or 56.0,
        )
        header_text = "\n".join(p.text for p in section.header.paragraphs if p.text.strip())
        footer_text = "\n".join(p.text for p in section.footer.paragraphs if p.text.strip())
        info.header = header_text.strip()[:200]
        info.footer = footer_text.strip()[:200]
        footer_xml = section.footer._element.xml if section.footer is not None else ""
        info.has_page_field = "PAGE" in footer_xml or "NUMPAGES" in footer_xml
    except Exception as exc:  # pragma: no cover - malformed sections
        info.warnings.append("Page setup could not be read; A4 defaults were used.")
        log.info("Section parsing issue: %s", exc)

    assets: dict[str, bytes] = {}
    body: list[str] = []
    list_open: str | None = None
    pending_break = False

    def close_list() -> None:
        nonlocal list_open
        if list_open:
            body.append(f"</{list_open}>")
            list_open = None

    for block in _iter_blocks(document):
        if hasattr(block, "rows"):
            close_list()
            if pending_break:
                body.append('<p style="page-break-before: always;"></p>')
                pending_break = False
            body.append(_table_html(block, assets))
            continue
        kind, _ = _style_kind(block)
        content = "".join(_run_html(run, assets) for run in block.runs)
        if not content.strip() and not block.text.strip():
            continue
        if block.paragraph_format.page_break_before:
            pending_break = True
        if pending_break:
            body.append('<p style="page-break-before: always;"></p>')
            pending_break = False
        inline_break = '<span class="page-break"></span>' in content
        if inline_break:
            content = content.replace('<span class="page-break"></span>', "")
        if kind.startswith("h"):
            close_list()
            body.append(f"<{kind}>{content}</{kind}>")
        elif kind == "li-bullet":
            if list_open != "ul":
                close_list()
                body.append("<ul>")
                list_open = "ul"
            body.append(f"<li>{content}</li>")
        elif kind == "li-number":
            if list_open != "ol":
                close_list()
                body.append("<ol>")
                list_open = "ol"
            body.append(f"<li>{content}</li>")
        elif kind == "blockquote":
            close_list()
            body.append(f"<blockquote>{content}</blockquote>")
        else:
            close_list()
            alignment = ""
            try:
                align = block.alignment
                if align is not None and int(align) == 1:
                    alignment = "text-align:center;"
                elif align is not None and int(align) == 2:
                    alignment = "text-align:right;"
                elif align is not None and int(align) == 3:
                    alignment = "text-align:justify;"
            except Exception:
                pass
            style_attr = f' style="{alignment}"' if alignment else ""
            body.append(f"<p{style_attr}>{content}</p>")
        if inline_break:
            body.append('<p style="page-break-before: always;"></p>')
    close_list()

    if not body:
        raise ConversionFailed("This Word document contains no renderable content.")

    css_font = "10.5pt"
    try:
        normal = document.styles["Normal"]
        if normal.font.size:
            css_font = f"{normal.font.size.pt:.1f}pt"
    except Exception:  # pragma: no cover
        pass

    html = f"<html><body>{''.join(body)}</body></html>"
    info.warnings.append(f"base-font:{css_font}")
    return html, assets, info


_CSS_TEMPLATE = """
body {{ font-family: 'DejaVu Sans', 'Liberation Sans', sans-serif; font-size: {font}; line-height: 1.45; color: #111827; }}
h1 {{ font-size: 1.9em; color: #0f172a; margin: 0 0 8pt 0; }}
h2 {{ font-size: 1.45em; color: #0f172a; margin: 12pt 0 6pt 0; }}
h3 {{ font-size: 1.2em; color: #1e293b; margin: 10pt 0 5pt 0; }}
h4, h5, h6 {{ font-size: 1.05em; color: #1e293b; margin: 9pt 0 4pt 0; }}
p {{ margin: 0 0 7pt 0; }}
ul, ol {{ margin: 0 0 8pt 16pt; padding: 0; }}
li {{ margin: 0 0 3pt 0; }}
table {{ border-collapse: collapse; width: 100%; margin: 8pt 0; }}
td, th {{ border: 0.6pt solid #cbd5e1; padding: 4pt 6pt; vertical-align: top; }}
th {{ background: #f1f5f9; text-align: left; }}
img {{ max-width: 100%; }}
blockquote {{ margin: 8pt 16pt; color: #475569; border-left: 2pt solid #cbd5e1; padding-left: 8pt; }}
sup, sub {{ font-size: 0.7em; }}
"""


def render(source: Path, destination: Path, workdir: Path) -> RenderInfo:
    html, assets, info = build_html(source)
    font_size = next((w.split(":", 1)[1] for w in info.warnings if w.startswith("base-font:")), "10.5pt")
    info.warnings = [w for w in info.warnings if not w.startswith("base-font:")]
    css = _CSS_TEMPLATE.format(font=font_size)

    if assets:
        assets_dir = workdir / "docx-assets"
        assets_dir.mkdir(parents=True, exist_ok=True)
        for name, blob in assets.items():
            (assets_dir / name).write_bytes(blob)
        story = pymupdf.Story(html=html, user_css=css, archive=str(assets_dir))
    else:
        story = pymupdf.Story(html=html, user_css=css)

    page_size = pymupdf.Rect(0, 0, info.page_width, info.page_height)
    top, bottom, left, right = info.margins
    content = pymupdf.Rect(left, top, info.page_width - right, info.page_height - bottom)
    if content.width < 100 or content.height < 100:
        content = pymupdf.Rect(56, 56, info.page_width - 56, info.page_height - 56)
    writer = pymupdf.DocumentWriter(str(destination))
    more = 1
    pages = 0
    try:
        while more:
            device = writer.begin_page(page_size)
            more, _ = story.place(content)
            story.draw(device)
            writer.end_page()
            pages += 1
            if pages > 2000:
                break
    finally:
        writer.close()
    if pages == 0:
        raise ConversionFailed("The Word document could not be laid out.")
    info.pages = pages
    if info.header or info.footer or info.has_page_field:
        _stamp_furniture(destination, info)
    v.validate_pdf(destination)
    return info


def _stamp_furniture(path: Path, info: RenderInfo) -> None:
    """Draw header/footer text and page numbers onto every rendered page."""
    try:
        doc = pymupdf.open(path)
    except Exception:  # pragma: no cover
        return
    try:
        from .fonts import page_font_alias

        total = doc.page_count
        top, bottom, left, right = info.margins
        for index, page in enumerate(doc):
            alias, font = page_font_alias(page)
            width = page.rect.width
            height = page.rect.height
            color = (0.42, 0.45, 0.5)
            if info.header:
                size = 8.0
                page.insert_text(
                    pymupdf.Point(left, max(18.0, top * 0.55)),
                    _truncate(info.header.replace("\n", " "), font, size, width - left - right),
                    fontsize=size,
                    fontname=alias,
                    color=color,
                )
                page.draw_line(
                    pymupdf.Point(left, max(24.0, top * 0.55) + 6),
                    pymupdf.Point(width - right, max(24.0, top * 0.55) + 6),
                    color=(0.85, 0.87, 0.9),
                    width=0.5,
                )
            footer_text = info.footer.replace("\n", " ").strip()
            if footer_text or info.has_page_field:
                size = 8.0
                number = f"{index + 1} / {total}"
                text = footer_text
                if info.has_page_field:
                    text = f"{text} {number}".strip() if text else number
                text = _truncate(text, font, size, width - left - right)
                baseline = height - max(16.0, bottom * 0.35)
                page.insert_text(
                    pymupdf.Point(left, baseline),
                    text,
                    fontsize=size,
                    fontname=alias,
                    color=color,
                )
        tmp = path.with_suffix(".stamped.pdf")
        doc.save(tmp, garbage=3, deflate=True)
    finally:
        doc.close()
    try:
        tmp.replace(path)
    except OSError:  # pragma: no cover
        pass


def _truncate(text: str, font, size: float, max_width: float) -> str:
    if font.text_length(text, fontsize=size) <= max_width:
        return text
    trimmed = text
    while trimmed and font.text_length(trimmed + "…", fontsize=size) > max_width:
        trimmed = trimmed[:-1]
    return trimmed + "…"
